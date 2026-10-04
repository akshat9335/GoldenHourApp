import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

def create_enhanced_slide_6_image():
    """Update original_slide_6.png with current prototype links and save enhanced_slide_6.png."""
    orig_path = "original_slide_6.png"
    out_path = "enhanced_slide_6.png"
    if not os.path.exists(orig_path):
        print(f"Error: {orig_path} not found!")
        return

    im = Image.open(orig_path).copy()
    draw = ImageDraw.Draw(im)

    sx = 2667.0 / 960.0
    sy = 1500.0 / 540.0

    # Clean the old links area with pure white
    draw.rectangle([int(32 * sx), int(244 * sy), int(445 * sx), int(363 * sy)], fill=(255, 255, 255))

    font_bold = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", int(14.0 * sy))
    font_regular = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", int(11.8 * sy))
    font_title = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", int(14.5 * sy))

    label_color = (33, 90, 140)
    link_color = (2, 132, 199)

    # 1. App link:
    y_pt = 246.0
    draw.text((int(36.4 * sx), int(y_pt * sy)), "App link (Prototype APK):", fill=label_color, font=font_title)
    y_pt += 17.5
    draw.text((int(36.4 * sx), int(y_pt * sy)), "https://drive.google.com/file/d/1DjNVllfmDO0dMCmqSco1yLzNUFJ_s1Ie/", fill=link_color, font=font_regular)
    y_pt += 14.5
    draw.text((int(36.4 * sx), int(y_pt * sy)), "view?usp=drivesdk", fill=link_color, font=font_regular)

    # 2. App Manual & Guide link:
    y_pt += 18.0
    draw.text((int(36.4 * sx), int(y_pt * sy)), "App Manual & Guide link (PDF):", fill=label_color, font=font_title)
    y_pt += 17.5
    draw.text((int(36.4 * sx), int(y_pt * sy)), "https://drive.google.com/file/d/13T6n8MtCcO78936hAwHKXWWar8z5qT0l/", fill=link_color, font=font_regular)
    y_pt += 14.5
    draw.text((int(36.4 * sx), int(y_pt * sy)), "view?usp=drivesdk", fill=link_color, font=font_regular)

    # 3. Git link:
    y_pt += 18.0
    draw.text((int(36.4 * sx), int(y_pt * sy)), "Git link :", fill=label_color, font=font_title)
    draw.text((int(95.0 * sx), int(y_pt * sy)), "https://github.com/akshat9335/GoldenHourApp", fill=link_color, font=font_bold)

    im.save(out_path)
    print(f"Generated {out_path} with updated prototype and manual links.")


def build_pptx():
    create_enhanced_slide_6_image()

    prs = Presentation()
    # Set slide dimensions to widescreen 16:9 (960pt x 540pt)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slides_info = [
        ("Slide 1", "original_slide_1.png"),
        ("Slide 2 (Enhanced Solution Flowchart)", "enhanced_slide_2.png"),
        ("Slide 3", "original_slide_3.png"),
        ("Slide 4", "original_slide_4.png"),
        ("Slide 5", "original_slide_5.png"),
        ("Slide 6 (Research & References with Clickable Links)", "enhanced_slide_6.png"),
    ]

    for title, img_path in slides_info:
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)
        print(f"Added {title} to PPTX")

        # If Slide 6, overlay transparent clickable hyperlink shapes
        if "Slide 6" in title:
            clickable_links = [
                # (name, x, y, w, h, url)
                ("Smart India Hackathon", 32, 106, 260, 18, "https://www.sih.gov.in/"),
                ("Ayushman Bharat Digital Mission (ABDM)", 32, 123, 325, 18, "https://abdm.gov.in/"),
                ("MoHFW Healthcare Guidelines", 32, 141, 270, 19, "https://www.mohfw.gov.in/"),
                ("ABDM Digital Health Framework", 32, 191, 220, 18, "https://abdm.gov.in/"),
                ("Telemedicine Guidelines", 32, 208, 180, 18, "https://www.mohfw.gov.in/pdf/Telemedicine.pdf"),
                ("Comprehensive Primary Healthcare", 32, 225, 235, 18, "https://nhsrcindia.org/"),
                ("App link (Prototype APK)", 32, 246, 420, 50, "https://drive.google.com/file/d/1DjNVllfmDO0dMCmqSco1yLzNUFJ_s1Ie/view?usp=drivesdk"),
                ("App Manual & Guide link (PDF)", 32, 299, 420, 50, "https://drive.google.com/file/d/13T6n8MtCcO78936hAwHKXWWar8z5qT0l/view?usp=drivesdk"),
                ("Git link", 32, 349, 360, 20, "https://github.com/akshat9335/GoldenHourApp"),
            ]

            for link_name, x, y, w, h, url in clickable_links:
                shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
                shape.fill.background()
                shape.line.fill.background()
                shape.click_action.hyperlink.address = url
                print(f"   Overlayed clickable hyperlink: {link_name} -> {url}")

    out_pptx = "Golden_Hour_SIH26133_Presentation.pptx"
    art_pptx = r"C:\Users\Admin\.gemini\antigravity\brain\160a95f8-5e11-42e1-8655-2a7f95c1d2d8\Golden_Hour_SIH26133_Presentation.pptx"

    prs.save(out_pptx)
    prs.save(art_pptx)
    print(f"Successfully saved {out_pptx} and copied to artifacts!")


if __name__ == "__main__":
    build_pptx()
