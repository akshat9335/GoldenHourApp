import os
import fitz
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from lxml import etree
from PIL import Image, ImageFilter

def enhance_techstack_image():
    """Extract and upscale tech stack image to 2455x1390 with unsharp masking."""
    source_pptx = r"C:\Users\Admin\Downloads\SIH2026_LifeLineCoders_GoldenHour_Clickable_Links.pptx"
    prs = Presentation(source_pptx)
    slide3 = prs.slides[2]
    sh18 = slide3.shapes[18]
    
    with open("temp_raw_techstack.png", "wb") as f:
        f.write(sh18.image.blob)
        
    im = Image.open("temp_raw_techstack.png")
    # Upscale 2.5x with high-quality Lanczos resampling
    im_high = im.resize((2455, 1390), Image.Resampling.LANCZOS)
    # Apply crisp unsharp mask filter to sharpen text and diagram boundaries
    im_sharp = im_high.filter(ImageFilter.UnsharpMask(radius=2, percent=180, threshold=2))
    im_sharp.save("enhanced_techstack_sharp.png")
    print("Generated enhanced_techstack_sharp.png at (2455x1390) resolution.")


def build_final_pptx():
    source_pptx = r"C:\Users\Admin\Downloads\SIH2026_LifeLineCoders_GoldenHour_Clickable_Links.pptx"
    prs = Presentation(source_pptx)

    # 1. Update Slide 3 (Technical Approach) with sharp high-res image
    slide3 = prs.slides[2]
    sh18 = slide3.shapes[18]
    blip_elem = sh18._element.xpath('.//a:blip')
    rId = blip_elem[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    image_part = slide3.part.related_part(rId)
    with open("enhanced_techstack_sharp.png", "rb") as f:
        image_part._blob = f.read()
    print("Replaced Slide 3 Tech Stack image in PPTX with high-res sharp version.")

    # 2. Update Slide 6 (Research & References) with full-coverage clickable solid-transparent shapes
    slide6 = prs.slides[5]

    clickable_links = [
        ("Smart India Hackathon 2026", 15.75, 96.0, 444.0, 45.0, "https://www.sih.gov.in/"),
        ("Ayushman Bharat Digital Mission (ABDM)", 15.75, 141.0, 444.0, 45.0, "https://abdm.gov.in/"),
        ("MoHFW Healthcare Guidelines", 15.75, 186.0, 444.0, 45.0, "https://www.mohfw.gov.in/"),
        ("Telemedicine Practice Guidelines", 15.75, 230.5, 444.0, 45.0, "https://esanjeevani.mohfw.gov.in/assets/guidelines/Telemedicine_Practice_Guidelines.pdf"),
        ("Git link", 15.75, 275.5, 444.0, 32.5, "https://github.com/akshat9335/GoldenHourApp.git"),
        ("App Manual & Guide link (PDF)", 15.75, 308.0, 444.0, 32.0, "https://drive.google.com/file/d/13T6n8MtCcO78936hAwHKXWWar8z5qT0l/view?usp=drivesdk"),
        ("App link (Prototype APK)", 15.75, 338.0, 444.0, 32.0, "https://drive.google.com/file/d/1DjNVllfmDO0dMCmqSco1yLzNUFJ_s1Ie/view?usp=drivesdk"),
    ]

    for name, x, y, w, h, url in clickable_links:
        shape = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)

        # 1% opacity = invisible to eye, solid to mouse clicks everywhere
        spPr = shape.element.spPr
        solidFill = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill')
        srgbClr = solidFill.find('{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
        alpha = etree.SubElement(srgbClr, '{http://schemas.openxmlformats.org/drawingml/2006/main}alpha')
        alpha.set('val', '1000')

        # No border line
        ln = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        if ln is not None:
            spPr.remove(ln)
        noFillLn = etree.SubElement(spPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}ln')
        etree.SubElement(noFillLn, '{http://schemas.openxmlformats.org/drawingml/2006/main}noFill')

        shape.click_action.hyperlink.address = url
        print(f"Added PPTX full-surface clickable hotspot: {name} -> {url}")

    out_pptx = "Golden_Hour_SIH26133_Presentation.pptx"
    art_pptx = r"C:\Users\Admin\.gemini\antigravity\brain\160a95f8-5e11-42e1-8655-2a7f95c1d2d8\Golden_Hour_SIH26133_Presentation.pptx"
    prs.save(out_pptx)
    prs.save(art_pptx)
    print(f"Successfully saved {out_pptx} and copied to artifacts!")


def build_final_pdf():
    pdf_path = r"C:\Users\Admin\.gemini\antigravity\brain\160a95f8-5e11-42e1-8655-2a7f95c1d2d8\.user_uploaded\media_1791143357210.pdf"
    doc = fitz.open(pdf_path)

    # 1. Update Page 3 (Technical Approach) with sharp high-res image
    p3 = doc[2]
    p3.replace_image(59, filename="enhanced_techstack_sharp.png")
    print("Replaced Page 3 image in PDF with high-res sharp version.")

    # 2. Update Page 6 (Research & References) with full-card links
    page6 = doc[5]
    for l in page6.get_links():
        page6.delete_link(l)

    clickable_links = [
        (fitz.Rect(15.75, 96.0, 459.75, 141.5), "https://www.sih.gov.in/"),
        (fitz.Rect(15.75, 141.5, 459.75, 186.0), "https://abdm.gov.in/"),
        (fitz.Rect(15.75, 186.0, 459.75, 230.5), "https://www.mohfw.gov.in/"),
        (fitz.Rect(15.75, 230.5, 459.75, 275.5), "https://esanjeevani.mohfw.gov.in/assets/guidelines/Telemedicine_Practice_Guidelines.pdf"),
        (fitz.Rect(15.75, 275.5, 459.75, 308.0), "https://github.com/akshat9335/GoldenHourApp.git"),
        (fitz.Rect(15.75, 308.0, 459.75, 338.0), "https://drive.google.com/file/d/13T6n8MtCcO78936hAwHKXWWar8z5qT0l/view?usp=drivesdk"),
        (fitz.Rect(15.75, 338.0, 459.75, 370.0), "https://drive.google.com/file/d/1DjNVllfmDO0dMCmqSco1yLzNUFJ_s1Ie/view?usp=drivesdk"),
    ]

    for rect, url in clickable_links:
        page6.insert_link({"kind": fitz.LINK_URI, "from": rect, "uri": url})
        print(f"Inserted PDF link across entire card: {rect} -> {url}")

    out_pdf = "Golden_Hour_SIH26133_Final_Presentation.pdf"
    art_pdf = r"C:\Users\Admin\.gemini\antigravity\brain\160a95f8-5e11-42e1-8655-2a7f95c1d2d8\Golden_Hour_SIH26133_Final_Presentation.pdf"
    doc.save(out_pdf)
    doc.save(art_pdf)
    print(f"Successfully saved {out_pdf} and copied to artifacts!")


if __name__ == "__main__":
    enhance_techstack_image()
    build_final_pptx()
    build_final_pdf()
